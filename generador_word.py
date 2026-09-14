import os
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# CONFIGURACIÓN GENERAL
# ============================================================

MAXIMO_SUFICIENCIAS = 7

# Estos atributos permiten que app.py consulte, sin romper la
# compatibilidad actual, si se creó un documento adicional.
ULTIMO_REPORTE_ESPECIAL = None
ULTIMO_RESUMEN_ESCRITURA = {}
ULTIMA_CONSTANCIA_APROBACION = None



# ============================================================
# COMPETENCIAS DEL FORMATO WORD (EDUCACIÓN)
# ============================================================

def _limpiar_texto_competencia(texto):
    """
    Limpia saltos de línea y espacios visuales de las celdas
    combinadas del formato de Educación.
    """
    return re.sub(
        r"\s+",
        " ",
        str(texto or "")
    ).strip()


def extraer_mapa_competencias_desde_formato(ruta_formato):
    """
    Lee directamente la proforma Word y devuelve:
        CURSO UPRIT NORMALIZADO -> COMPETENCIA

    Está pensado especialmente para los formatos de Educación,
    donde la columna COMPETENCIAS utiliza celdas combinadas
    verticalmente.

    Si el formato no contiene una tabla de competencias, devuelve {}.
    """
    if (
        not ruta_formato
        or not os.path.exists(ruta_formato)
    ):
        return {}

    documento = Document(ruta_formato)

    for tabla in documento.tables:

        if not tabla.rows:
            continue

        fila_encabezado = None
        indices_competencia = []
        idx_asignatura = None

        # Buscar encabezado en las primeras filas.
        for fila_idx, fila in enumerate(tabla.rows[:6]):

            textos = [
                normalizar(celda.text)
                for celda in fila.cells
            ]

            comp_indices = [
                i
                for i, texto in enumerate(textos)
                if texto == "COMPETENCIAS"
                or texto == "COMPETENCIA"
            ]

            asignatura_indices = [
                i
                for i, texto in enumerate(textos)
                if (
                    texto == "ASIGNATURA"
                    or texto == "ASIGNATURA UPRIT"
                )
            ]

            # Evitar tomar ASIGNATURA CONVALIDANTE.
            asignatura_indices = [
                i
                for i in asignatura_indices
                if "CONVALIDANTE" not in textos[i]
            ]

            if comp_indices and asignatura_indices:
                fila_encabezado = fila_idx
                indices_competencia = comp_indices
                idx_asignatura = asignatura_indices[-1]
                break

        if fila_encabezado is None:
            continue

        # En los formatos de Educación el encabezado COMPETENCIAS
        # suele abarcar dos columnas:
        #   1) tipo general/específica
        #   2) competencia real (COMUNICACIÓN, USO INFORMACIÓN, etc.)
        # La competencia académica es la columna más a la derecha.
        idx_competencia = max(indices_competencia)

        mapa = {}
        competencia_anterior = ""

        for fila in tabla.rows[fila_encabezado + 1:]:

            if (
                idx_asignatura >= len(fila.cells)
                or idx_competencia >= len(fila.cells)
            ):
                continue

            curso = _limpiar_texto_competencia(
                fila.cells[idx_asignatura].text
            )

            curso_norm = normalizar(curso)

            if (
                not curso_norm
                or "TOTAL DE CREDITOS" in curso_norm
                or curso_norm in ("ASIGNATURA", "ASIGNATURAS")
            ):
                continue

            competencia = _limpiar_texto_competencia(
                fila.cells[idx_competencia].text
            )

            competencia_norm = normalizar(competencia)

            # Ignorar la clasificación macro.
            if competencia_norm in (
                "COMPETENCIAS GENERALES",
                "COMPETENCIAS ESPECIFICAS",
                "COMPETENCIAS ESPECÍFICAS"
            ):
                competencia = ""

            # Algunas variantes de Word dejan vacías las filas
            # interiores de una celda combinada. En ese caso se
            # propaga la última competencia válida.
            if competencia:
                competencia_anterior = competencia
            else:
                competencia = competencia_anterior

            if competencia:
                mapa[curso_norm] = competencia

        if mapa:
            return mapa

    return {}


def encontrar_tabla_resumen_modalidad_educacion(documento):
    """
    Localiza la tabla final de Educación:
    CÓDIGO | ASIGNATURA UPRIT | CRÉDITOS |
    ASIGNATURA CONVALIDANTE | NOTA CONVALIDANTE |
    MODALIDAD DE CONVALIDACIÓN.
    """
    for tabla in documento.tables:

        texto = normalizar(
            " ".join(
                celda.text
                for fila in tabla.rows[:5]
                for celda in fila.cells
            )
        )

        if (
            "ASIGNATURA UPRIT" in texto
            and "MODALIDAD DE CONVALIDACION" in texto
            and "ASIGNATURA CONVALIDANTE" in texto
        ):
            return tabla

    return None


def rellenar_resumen_modalidad_educacion(
    documento,
    convalidaciones,
    suficiencias
):
    """
    Actualiza el cuadro final de los formatos de Educación.
    Si el documento no contiene ese cuadro, no hace nada.
    """
    tabla = encontrar_tabla_resumen_modalidad_educacion(
        documento
    )

    if tabla is None or not tabla.rows:
        return

    encabezados = [
        normalizar(celda.text)
        for celda in tabla.rows[0].cells
    ]

    def buscar_columna(*terminos):
        for i, texto in enumerate(encabezados):
            if all(
                termino in texto
                for termino in terminos
            ):
                return i
        return None

    idx_asignatura = buscar_columna(
        "ASIGNATURA UPRIT"
    )
    idx_convalidante = buscar_columna(
        "ASIGNATURA CONVALIDANTE"
    )
    idx_nota = buscar_columna(
        "NOTA CONVALIDANTE"
    )
    idx_modalidad = buscar_columna(
        "MODALIDAD",
        "CONVALIDACION"
    )

    if None in (
        idx_asignatura,
        idx_convalidante,
        idx_nota,
        idx_modalidad
    ):
        return

    mapa_resultados = {}

    for item in convalidaciones or []:
        if not isinstance(item, dict):
            continue

        destino = normalizar(
            item.get(
                "curso_destino",
                ""
            )
        )

        if not destino:
            continue

        existente = mapa_resultados.get(
            destino
        )

        if (
            existente is None
            or (
                valor_vacio(
                    existente.get(
                        "curso_origen"
                    )
                )
                and not valor_vacio(
                    item.get(
                        "curso_origen"
                    )
                )
            )
        ):
            mapa_resultados[destino] = item

    suficiencias_norm = {
        normalizar(curso)
        for curso in (suficiencias or [])
        if str(curso or "").strip()
    }

    for fila in tabla.rows[1:]:

        cells = fila.cells

        max_idx = max(
            idx_asignatura,
            idx_convalidante,
            idx_nota,
            idx_modalidad
        )

        if len(cells) <= max_idx:
            continue

        curso_uprit = cells[
            idx_asignatura
        ].text.strip()

        curso_norm = normalizar(
            curso_uprit
        )

        if (
            not curso_norm
            or "TOTAL" in curso_norm
        ):
            continue

        item = mapa_resultados.get(
            curso_norm,
            {}
        )

        if curso_norm in suficiencias_norm:

            limpiar_celda(
                cells[idx_convalidante]
            )
            limpiar_celda(
                cells[idx_nota]
            )
            poner_texto(
                cells[idx_modalidad],
                "EXAMEN DE SUFICIENCIA"
            )

            continue

        curso_origen = item.get(
            "curso_origen",
            ""
        )

        nota_final = item.get(
            "nota_convalidante",
            ""
        )

        if not valor_vacio(
            curso_origen
        ):
            poner_texto(
                cells[idx_convalidante],
                curso_origen
            )
        else:
            limpiar_celda(
                cells[idx_convalidante]
            )

        if not valor_vacio(
            nota_final
        ):
            poner_texto(
                cells[idx_nota],
                formatear_numero(
                    nota_final
                )
            )
        else:
            limpiar_celda(
                cells[idx_nota]
            )

        poner_texto(
            cells[idx_modalidad],
            "POR COMPETENCIA"
        )


# ============================================================
# FUNCIÓN PRINCIPAL
# ============================================================

def generar_word(
    datos_alumno,
    convalidaciones,
    suficiencias=None,
    ruta_formato="formato.docx",
    competencias=None,
    caso_especial=False,
    observacion_caso_especial="",
    suficiencias_iniciales=None,
    suficiencias_post_revision_completas=None,
    suficiencias_excedentes_revision=None,
    recomendaciones_caso_especial=None,
    aprobacion_coordinador=None,
    firma_coordinador_path=None
):
    """
    Genera la proforma Word de convalidación.

    REGLAS DE ESCRITURA:

    1. ASIGNATURA CONVALIDANTE:
       se copia exclusivamente desde curso_origen.

    2. NOTA PARCIAL:
       se copia exclusivamente desde item["nota"].
       Esa nota debe proceder del certificado.

    3. NOTA CONVALIDANTE:
       se copia exclusivamente desde item["nota_convalidante"],
       calculada previamente por convalidaciones.py.

    4. Si una competencia tiene un solo curso faltante:
       - curso_origen queda vacío;
       - nota parcial queda vacía;
       - nota convalidante de la competencia sí se conserva.

    5. Si la competencia tiene dos o más faltantes:
       - las equivalencias encontradas siguen llenas;
       - la nota convalidante queda vacía;
       - los faltantes pasan a suficiencia.

    6. Cuando convalidaciones.py marca CASO ESPECIAL:
       - la proforma principal SE GENERA;
       - la tabla principal muestra como máximo 7 suficiencias;
       - se crea además CASO_ESPECIAL_<ALUMNO>.docx;
       - el reporte registra la revisión académica especial.
    """

    global ULTIMO_REPORTE_ESPECIAL
    global ULTIMO_RESUMEN_ESCRITURA
    global ULTIMA_CONSTANCIA_APROBACION

    ULTIMO_REPORTE_ESPECIAL = None
    ULTIMO_RESUMEN_ESCRITURA = {}
    ULTIMA_CONSTANCIA_APROBACION = None

    if suficiencias is None:
        suficiencias = []

    if suficiencias_iniciales is None:
        suficiencias_iniciales = list(
            suficiencias
        )

    if suficiencias_post_revision_completas is None:
        suficiencias_post_revision_completas = list(
            suficiencias
        )

    if suficiencias_excedentes_revision is None:
        suficiencias_excedentes_revision = []

    if competencias is None:
        competencias = []

    if recomendaciones_caso_especial is None:
        recomendaciones_caso_especial = []

    if aprobacion_coordinador is None:
        aprobacion_coordinador = {}

    if convalidaciones is None:
        convalidaciones = []

    if not ruta_formato:
        ruta_formato = "formato.docx"

    if not os.path.exists(
        ruta_formato
    ):
        raise FileNotFoundError(
            f"No se encontró la proforma Word: {ruta_formato}"
        )

    documento = Document(
        ruta_formato
    )

    # --------------------------------------------------------
    # 1. DATOS GENERALES
    # --------------------------------------------------------

    rellenar_datos_generales(
        documento,
        datos_alumno
    )

    # --------------------------------------------------------
    # 2. CONVALIDACIONES
    # --------------------------------------------------------

    resumen_escritura = rellenar_convalidaciones(
        documento=documento,
        convalidaciones=convalidaciones
    )

    ULTIMO_RESUMEN_ESCRITURA = dict(
        resumen_escritura
    )

    # --------------------------------------------------------
    # 3. SUFICIENCIAS
    # --------------------------------------------------------

    rellenar_suficiencias(
        documento=documento,
        suficiencias=suficiencias,
        caso_especial=caso_especial
    )

    # --------------------------------------------------------
    # 3.1 RESUMEN FINAL DE MODALIDAD (FORMATOS EDUCACIÓN)
    # --------------------------------------------------------

    rellenar_resumen_modalidad_educacion(
        documento=documento,
        convalidaciones=convalidaciones,
        suficiencias=suficiencias
    )

    # --------------------------------------------------------
    # 4. GUARDAR PROFORMA PRINCIPAL
    # --------------------------------------------------------

    os.makedirs(
        "resultados",
        exist_ok=True
    )

    nombre = datos_alumno.get(
        "nombre",
        "ALUMNO"
    )

    nombre_archivo = limpiar_nombre_archivo(
        nombre
    )

    carrera_destino = limpiar_nombre_archivo(
        datos_alumno.get(
            "carrera_destino",
            ""
        )
    )

    if carrera_destino:
        nombre_salida = (
            f"CONVALIDACION_{carrera_destino}_{nombre_archivo}.docx"
        )
    else:
        nombre_salida = (
            f"CONVALIDACION_{nombre_archivo}.docx"
        )

    ruta = os.path.join(
        "resultados",
        nombre_salida
    )

    documento.save(
        ruta
    )

    # --------------------------------------------------------
    # 4.1 CONSTANCIA SEPARADA DE REVISIÓN/APROBACIÓN
    # --------------------------------------------------------

    if (
        isinstance(
            aprobacion_coordinador,
            dict
        )
        and aprobacion_coordinador.get(
            "aprobado"
        )
    ):
        ULTIMA_CONSTANCIA_APROBACION = (
            generar_constancia_aprobacion(
                datos_alumno=datos_alumno,
                aprobacion=aprobacion_coordinador,
                nombre_archivo=nombre_archivo
            )
        )

    # --------------------------------------------------------
    # 5. CASO ESPECIAL
    # --------------------------------------------------------

    if (
        caso_especial
        or len(
            suficiencias
        ) > MAXIMO_SUFICIENCIAS
    ):

        ULTIMO_REPORTE_ESPECIAL = generar_reporte_caso_especial(
            datos_alumno=datos_alumno,
            competencias=competencias,
            suficiencias=suficiencias,
            observacion=observacion_caso_especial,
            nombre_archivo=nombre_archivo,
            suficiencias_iniciales=suficiencias_iniciales,
            suficiencias_post_revision_completas=(
                suficiencias_post_revision_completas
            ),
            suficiencias_excedentes_revision=(
                suficiencias_excedentes_revision
            ),
            recomendaciones_caso_especial=(
                recomendaciones_caso_especial
            ),
            aprobacion_coordinador=(
                aprobacion_coordinador
            ),
            firma_coordinador_path=(
                firma_coordinador_path
            )
        )

    # Mantener compatibilidad con app.py:
    # la función continúa devolviendo SOLO la ruta principal.
    generar_word.ultimo_reporte_especial = (
        ULTIMO_REPORTE_ESPECIAL
    )

    generar_word.ultimo_resumen_escritura = (
        ULTIMO_RESUMEN_ESCRITURA
    )

    generar_word.ultima_constancia_aprobacion = (
        ULTIMA_CONSTANCIA_APROBACION
    )

    return ruta


# ============================================================
# DATOS GENERALES
# ============================================================

def rellenar_datos_generales(
    documento,
    datos
):
    """
    Rellena la primera tabla de datos generales.
    Conserva compatibilidad con las proformas existentes.
    """

    if not documento.tables:
        raise ValueError(
            "La proforma no contiene tablas."
        )

    # Buscar la tabla real de datos del estudiante.
    # En varios formatos de Educación, documento.tables[0] es el
    # encabezado institucional y la ficha del alumno está en tables[1].
    tabla = None

    for candidata in documento.tables:
        texto_candidata = normalizar(
            " ".join(
                celda.text
                for fila in candidata.rows[:8]
                for celda in fila.cells
            )
        )

        if (
            "APELLIDOS Y NOMBRES" in texto_candidata
            and "DNI" in texto_candidata
            and (
                "CORREO INSTITUCIONAL" in texto_candidata
                or "CORREO" in texto_candidata
            )
        ):
            tabla = candidata
            break

    if tabla is None:
        # Compatibilidad con formatos antiguos.
        tabla = documento.tables[0]

    # Apellidos y nombres
    if (
        len(tabla.rows) > 0
        and len(tabla.rows[0].cells) > 1
    ):
        poner_texto(
            tabla.rows[0].cells[1],
            datos.get(
                "nombre",
                ""
            )
        )

    # DNI y teléfono
    if len(tabla.rows) > 1:

        if len(tabla.rows[1].cells) > 1:
            poner_texto(
                tabla.rows[1].cells[1],
                datos.get(
                    "dni",
                    ""
                )
            )

        if len(tabla.rows[1].cells) > 3:
            poner_texto(
                tabla.rows[1].cells[3],
                datos.get(
                    "telefono",
                    ""
                )
            )

    # Email
    if (
        len(tabla.rows) > 2
        and len(tabla.rows[2].cells) > 1
    ):
        poner_texto(
            tabla.rows[2].cells[1],
            datos.get(
                "email",
                ""
            )
        )

    # Tipo de convalidación
    if (
        len(tabla.rows) > 3
        and len(tabla.rows[3].cells) > 1
    ):
        poner_texto(
            tabla.rows[3].cells[1],
            "CONVALIDACIÓN POR COMPETENCIAS"
        )

    # Institución de procedencia
    if (
        len(tabla.rows) > 4
        and len(tabla.rows[4].cells) > 1
    ):
        poner_texto(
            tabla.rows[4].cells[1],
            datos.get(
                "institucion",
                ""
            )
        )

    # Carrera de procedencia
    if (
        len(tabla.rows) > 5
        and len(tabla.rows[5].cells) > 1
    ):

        carrera_procedencia = (
            datos.get(
                "carrera_procedencia",
                ""
            )
            or datos.get(
                "carrera",
                ""
            )
        )

        poner_texto(
            tabla.rows[5].cells[1],
            carrera_procedencia
        )


# ============================================================
# DETECCIÓN ROBUSTA DE TABLA Y COLUMNAS
# ============================================================

def detectar_columnas_convalidaciones(
    tabla
):
    """
    Detecta las columnas aunque:
    - haya celdas combinadas;
    - los encabezados estén partidos;
    - existan variantes ortográficas como ASIGANTURA;
    - la cantidad de columnas cambie entre carreras.
    """

    if not tabla.rows:
        return None

    filas_revision = min(
        7,
        len(tabla.rows)
    )

    max_columnas = max(
        len(tabla.rows[i].cells)
        for i in range(filas_revision)
    )

    textos_columnas = []

    for col in range(
        max_columnas
    ):

        partes = []

        for fila_idx in range(
            filas_revision
        ):

            fila = tabla.rows[
                fila_idx
            ]

            if col < len(fila.cells):

                texto = normalizar(
                    fila.cells[col].text
                )

                if texto:
                    partes.append(
                        texto
                    )

        textos_columnas.append(
            " ".join(partes)
        )

    columnas = {
        "fila_encabezado": 0,
        "asignatura": None,
        "convalidante": None,
        "nota_parcial": None,
        "nota_convalidante": None
    }

    # --------------------------------------------------------
    # PRIMER INTENTO: TEXTO ACUMULADO POR COLUMNA
    # --------------------------------------------------------

    for i, texto in enumerate(
        textos_columnas
    ):

        if (
            "ASIGNATURA CONVALIDANTE" in texto
            or "ASIGANTURA CONVALIDANTE" in texto
            or (
                "CONVALIDANTE" in texto
                and "NOTA" not in texto
            )
        ):
            columnas[
                "convalidante"
            ] = i
            continue

        if (
            "NOTA PARCIAL" in texto
            or (
                "NOTA" in texto
                and "PARCIAL" in texto
            )
        ):
            columnas[
                "nota_parcial"
            ] = i
            continue

        if (
            "NOTA CONVALIDANTE" in texto
            or "NOTA CONVALIDANT" in texto
            or (
                "NOTA" in texto
                and (
                    "CONVALIDANTE" in texto
                    or "CONVALIDANT" in texto
                )
            )
        ):
            columnas[
                "nota_convalidante"
            ] = i
            continue

        if (
            "CONVALIDANTE" not in texto
            and (
                texto == "ASIGNATURA"
                or texto == "ASIGNATURAS"
                or "ASIGNATURA UPRIT" in texto
                or "ASIGNATURA A CONVALIDAR" in texto
            )
        ):
            columnas[
                "asignatura"
            ] = i

    # --------------------------------------------------------
    # SEGUNDO INTENTO: CELDA POR CELDA
    # --------------------------------------------------------

    for fila_idx in range(
        filas_revision
    ):

        fila = tabla.rows[
            fila_idx
        ]

        for i, celda in enumerate(
            fila.cells
        ):

            texto = normalizar(
                celda.text
            )

            if not texto:
                continue

            if (
                columnas["convalidante"] is None
                and (
                    "ASIGNATURA CONVALIDANTE" in texto
                    or "ASIGANTURA CONVALIDANTE" in texto
                )
            ):
                columnas[
                    "convalidante"
                ] = i

            elif (
                columnas["nota_parcial"] is None
                and "NOTA PARCIAL" in texto
            ):
                columnas[
                    "nota_parcial"
                ] = i

            elif (
                columnas["nota_convalidante"] is None
                and (
                    "NOTA CONVALIDANTE" in texto
                    or "NOTA CONVALIDANT" in texto
                )
            ):
                columnas[
                    "nota_convalidante"
                ] = i

            elif (
                columnas["asignatura"] is None
                and "CONVALIDANTE" not in texto
                and (
                    texto == "ASIGNATURA"
                    or texto == "ASIGNATURAS"
                    or "ASIGNATURA UPRIT" in texto
                )
            ):
                columnas[
                    "asignatura"
                ] = i

                columnas[
                    "fila_encabezado"
                ] = fila_idx

    if all(
        columnas[clave] is not None
        for clave in [
            "asignatura",
            "convalidante",
            "nota_parcial",
            "nota_convalidante"
        ]
    ):
        return columnas

    return None


def encontrar_tabla_convalidaciones(
    documento
):
    """
    Devuelve la tabla principal y las columnas detectadas.
    """

    candidatos = []

    for indice, tabla in enumerate(
        documento.tables
    ):

        columnas = detectar_columnas_convalidaciones(
            tabla
        )

        if columnas:

            # Preferir tabla que además contenga términos propios
            # de la proforma de convalidaciones.
            texto_tabla = normalizar(
                " ".join(
                    celda.text
                    for fila in tabla.rows[:10]
                    for celda in fila.cells
                )
            )

            puntaje = 0

            if "CONVALID" in texto_tabla:
                puntaje += 2

            if "CREDITO" in texto_tabla:
                puntaje += 1

            if "CODIGO" in texto_tabla:
                puntaje += 1

            candidatos.append(
                (
                    puntaje,
                    indice,
                    tabla,
                    columnas
                )
            )

    if not candidatos:
        raise ValueError(
            "No se encontró la tabla principal de convalidaciones. "
            "No fue posible identificar las columnas ASIGNATURA, "
            "ASIGNATURA CONVALIDANTE, NOTA PARCIAL y "
            "NOTA CONVALIDANTE."
        )

    candidatos.sort(
        key=lambda x: x[0],
        reverse=True
    )

    _, _, tabla, columnas = candidatos[
        0
    ]

    return tabla, columnas



# ============================================================
# MAPA SEGURO DE NOTAS POR CURSO DE PROCEDENCIA
# ============================================================

def construir_mapa_notas_origen(
    convalidaciones
):
    """
    Construye un mapa:
        CURSO ORIGEN NORMALIZADO -> NOTA REAL

    Solo usa notas explícitamente presentes en los propios resultados
    de convalidaciones.py.

    Nunca relaciona una nota con un curso distinto.
    Si un mismo curso aparece con notas diferentes, no se usa como
    respaldo automático para evitar ambigüedad.
    """

    candidatos = {}

    for item in convalidaciones or []:

        if not isinstance(
            item,
            dict
        ):
            continue

        curso_origen = str(
            item.get(
                "curso_origen",
                ""
            )
            or ""
        ).strip()

        if not curso_origen:
            continue

        nota = item.get(
            "nota",
            ""
        )

        if valor_vacio(
            nota
        ):
            continue

        nota_texto = formatear_numero(
            nota
        )

        if valor_vacio(
            nota_texto
        ):
            continue

        clave = normalizar(
            curso_origen
        )

        if not clave:
            continue

        candidatos.setdefault(
            clave,
            set()
        ).add(
            nota_texto
        )

    mapa = {}

    for clave, notas in candidatos.items():

        if len(
            notas
        ) == 1:
            mapa[
                clave
            ] = next(
                iter(
                    notas
                )
            )

    return mapa


# ============================================================
# RELLENAR TABLA PRINCIPAL
# ============================================================

def rellenar_convalidaciones(
    documento,
    convalidaciones
):
    """
    Copia al Word exactamente lo producido por convalidaciones.py.

    Compatible con:
    - formatos clásicos de una sola tabla;
    - formatos de Educación cuya tabla principal continúa en otra
      tabla de Word en la página siguiente;
    - encabezados truncados como NOTA CONVALIDANT.
    """

    tabla_principal, columnas = encontrar_tabla_convalidaciones(
        documento
    )

    idx_asignatura = columnas["asignatura"]
    idx_convalidante = columnas["convalidante"]
    idx_nota_parcial = columnas["nota_parcial"]
    idx_nota_convalidante = columnas["nota_convalidante"]

    inicio_principal = columnas["fila_encabezado"] + 1

    # --------------------------------------------------------
    # DETECTAR CONTINUACIONES DE LA TABLA PRINCIPAL
    # --------------------------------------------------------
    tablas_academicas = [
        (tabla_principal, inicio_principal)
    ]

    try:
        indice_principal = next(
            i
            for i, tabla in enumerate(documento.tables)
            if tabla._tbl is tabla_principal._tbl
        )
    except StopIteration:
        indice_principal = -1

    columnas_necesarias = max(
        idx_asignatura,
        idx_convalidante,
        idx_nota_parcial,
        idx_nota_convalidante
    ) + 1

    if indice_principal >= 0:
        for tabla in documento.tables[indice_principal + 1:]:

            texto_tabla = normalizar(
                " ".join(
                    celda.text
                    for fila in tabla.rows[:8]
                    for celda in fila.cells
                )
            )

            # Al llegar a suficiencias o no convalidadas termina
            # la continuidad de la proforma principal.
            if (
                "EXAMEN DE SUFICIENCIA" in texto_tabla
                or "ASIGNATURAS NO CONVALIDADAS" in texto_tabla
            ):
                break

            if not tabla.rows:
                continue

            max_cols = max(
                len(fila.cells)
                for fila in tabla.rows
            )

            if max_cols < columnas_necesarias:
                continue

            # La continuación debe tener contenido académico en la
            # misma columna ASIGNATURA usada por la tabla principal.
            tiene_asignaturas = False

            for fila in tabla.rows[:12]:
                if idx_asignatura >= len(fila.cells):
                    continue

                curso = normalizar(
                    fila.cells[idx_asignatura].text
                )

                if (
                    curso
                    and curso not in (
                        "ASIGNATURA",
                        "ASIGNATURAS"
                    )
                    and "TOTAL DE CREDITOS" not in curso
                ):
                    tiene_asignaturas = True
                    break

            if tiene_asignaturas:
                tablas_academicas.append(
                    (tabla, 0)
                )

    # --------------------------------------------------------
    # FILAS ACADÉMICAS DE TODAS LAS PARTES
    # --------------------------------------------------------
    filas_academicas = []

    for tabla_actual, inicio in tablas_academicas:

        for numero_fila, fila in enumerate(
            tabla_actual.rows[inicio:],
            start=inicio
        ):

            cells = fila.cells

            if len(cells) < columnas_necesarias:
                continue

            curso_word = cells[
                idx_asignatura
            ].text.strip()

            curso_word_normal = normalizar(
                curso_word
            )

            if not curso_word_normal:
                continue

            if (
                "TOTAL DE CREDITOS" in curso_word_normal
                or "TOTAL CREDITOS" in curso_word_normal
                or curso_word_normal in (
                    "ASIGNATURA",
                    "ASIGNATURAS"
                )
            ):
                continue

            filas_academicas.append({
                "numero_fila": numero_fila,
                "fila": fila,
                "curso": curso_word,
                "curso_normal": curso_word_normal
            })

    # --------------------------------------------------------
    # MAPA DE FILAS DEL WORD
    # --------------------------------------------------------
    filas_por_curso = {}

    for dato_fila in filas_academicas:
        curso_normal = dato_fila["curso_normal"]

        if curso_normal not in filas_por_curso:
            filas_por_curso[curso_normal] = dato_fila

    # --------------------------------------------------------
    # LIMPIAR SOLO COLUMNAS DE SALIDA
    # --------------------------------------------------------
    celdas_limpiadas = set()

    for dato_fila in filas_academicas:
        cells = dato_fila["fila"].cells

        for indice in [
            idx_convalidante,
            idx_nota_parcial,
            idx_nota_convalidante
        ]:
            celda = cells[indice]
            identidad = id(celda._tc)

            if identidad in celdas_limpiadas:
                continue

            limpiar_celda(celda)
            celdas_limpiadas.add(identidad)

    # --------------------------------------------------------
    # PREPARAR RESULTADOS
    # --------------------------------------------------------
    mapa_resultados = {}

    for item in convalidaciones or []:

        if not isinstance(item, dict):
            continue

        curso_destino = normalizar(
            item.get(
                "curso_destino",
                ""
            )
        )

        if not curso_destino:
            continue

        if curso_destino not in mapa_resultados:
            mapa_resultados[curso_destino] = item

    mapa_notas_origen = construir_mapa_notas_origen(
        convalidaciones
    )

    filas_reales_escritas = 0
    filas_competencia_vacias = 0
    cursos_no_encontrados = []
    notas_parciales_escritas = {}
    asignaturas_escritas = {}
    notas_recuperadas_mismo_origen = {}
    notas_convalidantes_por_celda = {}

    # --------------------------------------------------------
    # ESCRIBIR CURSO POR CURSO
    # --------------------------------------------------------
    for curso_destino_normal, item in mapa_resultados.items():

        dato_fila = filas_por_curso.get(
            curso_destino_normal
        )

        # Respaldo para pequeñas diferencias tipográficas.
        if dato_fila is None:

            mejor_dato = None
            mejor_puntaje = 0

            for candidato in filas_academicas:

                puntajes = fuzz_ratio_seguro(
                    curso_destino_normal,
                    candidato["curso_normal"]
                )

                puntaje = max(puntajes)

                if puntaje > mejor_puntaje:
                    mejor_puntaje = puntaje
                    mejor_dato = candidato

            if (
                mejor_dato is None
                or mejor_puntaje < 96
            ):
                cursos_no_encontrados.append(
                    item.get(
                        "curso_destino",
                        ""
                    )
                )
                continue

            dato_fila = mejor_dato

        fila = dato_fila["fila"]
        cells = fila.cells

        curso_origen = item.get(
            "curso_origen",
            ""
        )

        nota_parcial = item.get(
            "nota",
            ""
        )

        # Respaldo seguro: solo la nota del MISMO curso origen.
        if (
            not valor_vacio(curso_origen)
            and valor_vacio(nota_parcial)
        ):

            clave_origen = normalizar(
                curso_origen
            )

            nota_respaldo = mapa_notas_origen.get(
                clave_origen,
                ""
            )

            if not valor_vacio(nota_respaldo):

                nota_parcial = nota_respaldo

                notas_recuperadas_mismo_origen[
                    item.get(
                        "curso_destino",
                        ""
                    )
                ] = {
                    "curso_origen": curso_origen,
                    "nota": nota_respaldo
                }

        nota_convalidante = item.get(
            "nota_convalidante",
            ""
        )

        estado = normalizar(
            item.get(
                "estado",
                ""
            )
        )

        estado_competencia = normalizar(
            item.get(
                "estado_competencia",
                ""
            )
        )

        es_fila_faltante_competencia = (
            valor_vacio(curso_origen)
            and (
                "CONVALIDACION POR COMPETENCIA" in estado
                or "CURSO FALTANTE" in estado
                or "CAMPO VACIO" in estado
            )
        )

        if es_fila_faltante_competencia:

            limpiar_celda(
                cells[idx_convalidante]
            )

            limpiar_celda(
                cells[idx_nota_parcial]
            )

            filas_competencia_vacias += 1

        else:

            if not valor_vacio(curso_origen):

                poner_texto(
                    cells[idx_convalidante],
                    curso_origen
                )

                asignaturas_escritas[
                    item.get(
                        "curso_destino",
                        ""
                    )
                ] = curso_origen

            if not valor_vacio(nota_parcial):

                texto_nota = formatear_numero(
                    nota_parcial
                )

                poner_texto(
                    cells[idx_nota_parcial],
                    texto_nota
                )

                notas_parciales_escritas[
                    item.get(
                        "curso_destino",
                        ""
                    )
                ] = texto_nota

            else:
                limpiar_celda(
                    cells[idx_nota_parcial]
                )

            if not valor_vacio(curso_origen):
                filas_reales_escritas += 1

        # NOTA CONVALIDANTE
        celda_nota_final = cells[
            idx_nota_convalidante
        ]

        identidad_nota = id(
            celda_nota_final._tc
        )

        if (
            "SUFICIENCIA" not in estado_competencia
            and not valor_vacio(nota_convalidante)
        ):

            valor_final = formatear_numero(
                nota_convalidante
            )

            valor_anterior = notas_convalidantes_por_celda.get(
                identidad_nota
            )

            if valor_anterior is None:

                poner_texto(
                    celda_nota_final,
                    valor_final
                )

                notas_convalidantes_por_celda[
                    identidad_nota
                ] = valor_final

            elif valor_anterior == valor_final:
                pass

            else:
                # Nunca sobrescribir una celda fusionada con otra nota.
                pass

        elif "SUFICIENCIA" in estado_competencia:

            if identidad_nota not in notas_convalidantes_por_celda:
                limpiar_celda(
                    celda_nota_final
                )

    return {
        "filas_reales_escritas":
            filas_reales_escritas,
        "filas_competencia_vacias":
            filas_competencia_vacias,
        "cursos_no_encontrados":
            cursos_no_encontrados,
        "cantidad_filas_academicas":
            len(filas_academicas),
        "cantidad_tablas_academicas":
            len(tablas_academicas),
        "asignaturas_escritas":
            asignaturas_escritas,
        "notas_parciales_escritas":
            notas_parciales_escritas,
        "notas_recuperadas_mismo_origen":
            notas_recuperadas_mismo_origen,
        "cantidad_notas_convalidantes":
            len(notas_convalidantes_por_celda)
    }


# ============================================================
# TABLA DE SUFICIENCIA
# ============================================================

def encontrar_tabla_suficiencia(
    documento
):

    candidatos = []

    for indice, tabla in enumerate(
        documento.tables
    ):

        texto = normalizar(
            " ".join(
                celda.text
                for fila in tabla.rows
                for celda in fila.cells
            )
        )

        if "SUFICIENCIA" not in texto:
            continue

        puntaje = 0

        if "ASIGNATURA" in texto:
            puntaje += 2

        if "NOTA" in texto:
            puntaje += 1

        candidatos.append(
            (
                puntaje,
                indice,
                tabla
            )
        )

    if candidatos:

        candidatos.sort(
            key=lambda x: x[0],
            reverse=True
        )

        return candidatos[
            0
        ][2]

    if len(
        documento.tables
    ) > 3:
        return documento.tables[
            3
        ]

    raise ValueError(
        "No se encontró la tabla de examen de suficiencia."
    )


def rellenar_suficiencias(
    documento,
    suficiencias,
    caso_especial=False
):

    tabla = encontrar_tabla_suficiencia(
        documento
    )

    # Detectar fila de encabezado
    fila_encabezado = 0

    for i, fila in enumerate(
        tabla.rows[:5]
    ):

        texto = normalizar(
            " ".join(
                celda.text
                for celda in fila.cells
            )
        )

        if (
            "SUFICIENCIA" in texto
            or (
                "ASIGNATURA" in texto
                and "NOTA" in texto
            )
        ):
            fila_encabezado = i

    inicio = (
        fila_encabezado
        + 1
    )

    # Garantizar siete filas útiles
    while len(
        tabla.rows
    ) < (
        inicio
        + MAXIMO_SUFICIENCIAS
    ):
        tabla.add_row()

    # Limpiar
    for posicion in range(
        MAXIMO_SUFICIENCIAS
    ):

        fila = tabla.rows[
            inicio
            + posicion
        ]

        numero = (
            posicion
            + 1
        )

        if len(fila.cells) >= 1:
            poner_texto(
                fila.cells[0],
                numero
            )

        if len(fila.cells) >= 2:
            limpiar_celda(
                fila.cells[1]
            )

        if len(fila.cells) >= 3:
            limpiar_celda(
                fila.cells[2]
            )

    # CASO ESPECIAL:
    # La proforma principal SIEMPRE se genera.
    # En la tabla principal de suficiencia se colocan hasta 7 cursos.
    # Si existen más de 7, el listado completo aparecerá además
    # en el documento independiente de CASO ESPECIAL.
    for posicion, curso in enumerate(
        suficiencias
    ):

        if posicion >= MAXIMO_SUFICIENCIAS:
            break

        fila = tabla.rows[
            inicio
            + posicion
        ]

        if len(fila.cells) >= 1:
            poner_texto(
                fila.cells[0],
                posicion + 1
            )

        if len(fila.cells) >= 2:
            poner_texto(
                fila.cells[1],
                curso
            )

        if len(fila.cells) >= 3:
            limpiar_celda(
                fila.cells[2]
            )


# ============================================================
# CONSTANCIA SEPARADA DE REVISIÓN Y APROBACIÓN
# ============================================================

def generar_constancia_aprobacion(
    datos_alumno,
    aprobacion,
    nombre_archivo
):
    """
    Genera un documento Word INDEPENDIENTE que deja constancia
    de la revisión y aprobación del expediente.

    La proforma oficial no se modifica.
    """

    if not isinstance(
        aprobacion,
        dict
    ):
        return None

    if not aprobacion.get(
        "aprobado"
    ):
        return None

    documento = Document()

    titulo = documento.add_heading(
        "CONSTANCIA DE REVISIÓN Y APROBACIÓN",
        level=1
    )
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitulo = documento.add_paragraph(
        "Proceso de Convalidación Académica"
    )
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    institucion = documento.add_paragraph(
        "Universidad Privada de Trujillo"
    )
    institucion.alignment = WD_ALIGN_PARAGRAPH.CENTER

    documento.add_paragraph()

    tabla = documento.add_table(
        rows=6,
        cols=2
    )
    tabla.style = "Table Grid"

    datos_tabla = [
        (
            "Alumno",
            datos_alumno.get(
                "nombre",
                ""
            )
        ),
        (
            "DNI",
            datos_alumno.get(
                "dni",
                ""
            )
        ),
        (
            "Institución de procedencia",
            datos_alumno.get(
                "institucion",
                ""
            )
        ),
        (
            "Carrera de procedencia",
            datos_alumno.get(
                "carrera_procedencia",
                datos_alumno.get(
                    "carrera",
                    ""
                )
            )
        ),
        (
            "Carrera UPRIT de destino",
            datos_alumno.get(
                "carrera_destino",
                ""
            )
        ),
        (
            "Estado del expediente",
            "REVISADO Y APROBADO"
        )
    ]

    for indice, (
        etiqueta,
        valor
    ) in enumerate(
        datos_tabla
    ):
        tabla.rows[
            indice
        ].cells[0].text = str(
            etiqueta
        )
        tabla.rows[
            indice
        ].cells[1].text = str(
            valor
            or ""
        )

        for run in tabla.rows[
            indice
        ].cells[0].paragraphs[0].runs:
            run.bold = True

    documento.add_paragraph()

    p = documento.add_paragraph()

    p.add_run(
        "CONSTANCIA: "
    ).bold = True

    p.add_run(
        "El expediente de convalidación académica ha sido revisado "
        "por el coordinador de la carrera y cuenta con su aprobación "
        "para continuar con el flujo institucional correspondiente."
    )

    documento.add_paragraph(
        "Esta constancia es un documento de control interno y no "
        "reemplaza las firmas, vistos buenos o aprobaciones que ya "
        "corresponden en la proforma oficial de convalidación."
    )

    documento.add_heading(
        "Datos de la revisión",
        level=2
    )

    p = documento.add_paragraph()
    p.add_run(
        "Coordinador: "
    ).bold = True
    p.add_run(
        str(
            aprobacion.get(
                "nombre",
                ""
            )
        )
    )

    p = documento.add_paragraph()
    p.add_run(
        "Cargo: "
    ).bold = True
    p.add_run(
        str(
            aprobacion.get(
                "cargo",
                "Coordinador de carrera"
            )
        )
    )

    p = documento.add_paragraph()
    p.add_run(
        "Fecha y hora de aprobación: "
    ).bold = True
    p.add_run(
        str(
            aprobacion.get(
                "fecha",
                ""
            )
        )
    )

    estado = documento.add_paragraph()
    estado.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_estado = estado.add_run(
        "REVISADO Y APROBADO"
    )
    run_estado.bold = True

    os.makedirs(
        "resultados",
        exist_ok=True
    )

    carrera_destino = limpiar_nombre_archivo(
        datos_alumno.get(
            "carrera_destino",
            ""
        )
    )

    if carrera_destino:
        nombre_constancia = (
            f"CONSTANCIA_APROBACION_{carrera_destino}_"
            f"{nombre_archivo}.docx"
        )
    else:
        nombre_constancia = (
            f"CONSTANCIA_APROBACION_{nombre_archivo}.docx"
        )

    ruta_constancia = os.path.join(
        "resultados",
        nombre_constancia
    )

    documento.save(
        ruta_constancia
    )

    return ruta_constancia


# ============================================================
# REPORTE DE CASO ESPECIAL
# ============================================================

def generar_reporte_caso_especial(
    datos_alumno,
    competencias,
    suficiencias,
    observacion,
    nombre_archivo,
    suficiencias_iniciales=None,
    suficiencias_post_revision_completas=None,
    suficiencias_excedentes_revision=None,
    recomendaciones_caso_especial=None,
    aprobacion_coordinador=None,
    firma_coordinador_path=None
):

    documento = Document()

    if suficiencias_iniciales is None:
        suficiencias_iniciales = list(
            suficiencias
        )

    if suficiencias_post_revision_completas is None:
        suficiencias_post_revision_completas = list(
            suficiencias
        )

    if suficiencias_excedentes_revision is None:
        suficiencias_excedentes_revision = []

    if recomendaciones_caso_especial is None:
        recomendaciones_caso_especial = []

    if aprobacion_coordinador is None:
        aprobacion_coordinador = {}

    carrera_destino_texto = str(
        datos_alumno.get(
            "carrera_destino",
            ""
        )
        or ""
    ).upper()

    titulo_reporte = (
        "REPORTE DE CASO ESPECIAL - INGENIERÍA CIVIL"
        if "CIVIL" in carrera_destino_texto
        else "REPORTE DE CASO ESPECIAL DE CONVALIDACIÓN"
    )

    documento.add_heading(
        titulo_reporte,
        level=1
    )

    documento.add_paragraph(
        "Universidad Privada de Trujillo"
    )

    documento.add_paragraph(
        f"Alumno: {datos_alumno.get('nombre', '')}"
    )

    documento.add_paragraph(
        f"DNI: {datos_alumno.get('dni', '')}"
    )

    documento.add_paragraph(
        "Institución de procedencia: "
        f"{datos_alumno.get('institucion', '')}"
    )

    documento.add_paragraph(
        "Carrera de procedencia: "
        f"{datos_alumno.get('carrera_procedencia', datos_alumno.get('carrera', ''))}"
    )

    documento.add_paragraph(
        "Carrera UPRIT de destino: "
        f"{datos_alumno.get('carrera_destino', '')}"
    )

    documento.add_heading(
        "Motivo del caso especial",
        level=2
    )

    if observacion:

        documento.add_paragraph(
            observacion
        )

    else:

        documento.add_paragraph(
            f"Se identificaron {len(suficiencias)} asignaturas "
            f"seleccionadas para examen de suficiencia después "
            f"de la evaluación académica especial."
        )

    documento.add_paragraph(
        "La proforma principal ha sido generada igualmente. "
        "Este reporte deja constancia de la segunda revisión "
        "académica efectuada y de las asignaturas finalmente "
        "seleccionadas para examen de suficiencia."
    )

    documento.add_heading(
        "1. Candidatos detectados antes de la segunda revisión",
        level=2
    )

    tabla_inicial = documento.add_table(
        rows=1,
        cols=2
    )

    tabla_inicial.rows[0].cells[0].text = "N.°"
    tabla_inicial.rows[0].cells[1].text = "Asignatura"

    for numero, curso in enumerate(
        suficiencias_iniciales,
        start=1
    ):
        fila = tabla_inicial.add_row().cells
        fila[0].text = str(numero)
        fila[1].text = str(curso)

    documento.add_paragraph(
        f"Total inicial: {len(suficiencias_iniciales)}"
    )

    documento.add_heading(
        "2. Cursos faltantes después de la revisión especialista",
        level=2
    )

    tabla_post = documento.add_table(
        rows=1,
        cols=2
    )

    tabla_post.rows[0].cells[0].text = "N.°"
    tabla_post.rows[0].cells[1].text = "Asignatura"

    for numero, curso in enumerate(
        suficiencias_post_revision_completas,
        start=1
    ):
        fila = tabla_post.add_row().cells
        fila[0].text = str(numero)
        fila[1].text = str(curso)

    documento.add_paragraph(
        "Total después de revisión: "
        f"{len(suficiencias_post_revision_completas)}"
    )

    documento.add_heading(
        "3. Asignaturas seleccionadas para examen de suficiencia",
        level=2
    )

    tabla = documento.add_table(
        rows=1,
        cols=2
    )

    tabla.rows[0].cells[0].text = "N.°"
    tabla.rows[0].cells[1].text = "Asignatura"

    for numero, curso in enumerate(
        suficiencias,
        start=1
    ):
        fila = tabla.add_row().cells
        fila[0].text = str(numero)
        fila[1].text = str(curso)

    documento.add_paragraph(
        f"Total seleccionado para suficiencia: {len(suficiencias)}"
    )

    if suficiencias_excedentes_revision:

        documento.add_heading(
            "4. Pendientes adicionales de revisión académica",
            level=2
        )

        tabla_excedentes = documento.add_table(
            rows=1,
            cols=2
        )

        tabla_excedentes.rows[0].cells[0].text = "N.°"
        tabla_excedentes.rows[0].cells[1].text = "Asignatura"

        for numero, curso in enumerate(
            suficiencias_excedentes_revision,
            start=1
        ):
            fila = tabla_excedentes.add_row().cells
            fila[0].text = str(numero)
            fila[1].text = str(curso)

        documento.add_paragraph(
            "Estos cursos NO se consideran convalidados y tampoco "
            "se agregan a la tabla principal de suficiencias. "
            "Requieren revisión académica adicional."
        )

    documento.add_heading(
        "5. Recomendaciones para revisión del coordinador",
        level=2
    )

    documento.add_paragraph(
        "Las siguientes sugerencias son referenciales y NO constituyen "
        "una convalidación automática. El coordinador debe contrastarlas "
        "con el certificado y, cuando corresponda, con los sílabos."
    )

    if recomendaciones_caso_especial:

        tabla_rec = documento.add_table(
            rows=1,
            cols=6
        )

        encabezados = [
            "Curso UPRIT pendiente",
            "Curso sugerido del certificado",
            "Nota",
            "Afinidad",
            "Recomendación",
            "Justificación"
        ]

        for indice, encabezado in enumerate(
            encabezados
        ):
            tabla_rec.rows[0].cells[
                indice
            ].text = encabezado

        for item in recomendaciones_caso_especial:

            fila = tabla_rec.add_row().cells
            fila[0].text = str(item.get("curso_destino", ""))
            fila[1].text = str(item.get("curso_sugerido", ""))
            fila[2].text = str(item.get("nota", ""))
            fila[3].text = str(item.get("afinidad", ""))
            fila[4].text = str(item.get("recomendacion", ""))
            fila[5].text = str(item.get("justificacion", ""))

    else:

        documento.add_paragraph(
            "No se identificaron recomendaciones adicionales."
        )

    if competencias:

        documento.add_heading(
            "Resumen por competencias",
            level=2
        )

        tabla_comp = documento.add_table(
            rows=1,
            cols=5
        )

        cabecera = tabla_comp.rows[
            0
        ].cells

        cabecera[0].text = "Competencia"
        cabecera[1].text = "Total"
        cabecera[2].text = "Equivalencias"
        cabecera[3].text = "Faltantes"
        cabecera[4].text = "Estado"

        for item in competencias:

            fila = tabla_comp.add_row().cells

            fila[0].text = str(
                item.get(
                    "competencia",
                    ""
                )
            )

            fila[1].text = str(
                item.get(
                    "total_asignaturas",
                    ""
                )
            )

            fila[2].text = str(
                item.get(
                    "asignaturas_con_equivalencia",
                    ""
                )
            )

            fila[3].text = str(
                item.get(
                    "asignaturas_faltantes",
                    ""
                )
            )

            fila[4].text = str(
                item.get(
                    "estado",
                    ""
                )
            )

    documento.add_paragraph(
        "El expediente deberá ser revisado por los evaluadores "
        "académicos competentes antes de adoptar una decisión."
    )

    carrera_destino = limpiar_nombre_archivo(
        datos_alumno.get(
            "carrera_destino",
            ""
        )
    )

    if carrera_destino:
        nombre_reporte = (
            f"CASO_ESPECIAL_{carrera_destino}_{nombre_archivo}.docx"
        )
    else:
        nombre_reporte = (
            f"CASO_ESPECIAL_{nombre_archivo}.docx"
        )

    ruta_especial = os.path.join(
        "resultados",
        nombre_reporte
    )

    documento.save(
        ruta_especial
    )

    return ruta_especial


# ============================================================
# UTILIDADES
# ============================================================

def poner_texto(
    celda,
    valor
):
    """
    Reemplaza texto intentando conservar el formato del primer run.
    """

    if valor is None:
        valor = ""

    valor = str(
        valor
    )

    if not celda.paragraphs:
        celda.text = valor
        return

    parrafo = celda.paragraphs[
        0
    ]

    if parrafo.runs:

        parrafo.runs[
            0
        ].text = valor

        for run in parrafo.runs[
            1:
        ]:
            run.text = ""

    else:

        parrafo.add_run(
            valor
        )

    for parrafo_extra in celda.paragraphs[
        1:
    ]:

        for run in parrafo_extra.runs:
            run.text = ""


def limpiar_celda(
    celda
):
    poner_texto(
        celda,
        ""
    )


def valor_vacio(
    valor
):
    if valor is None:
        return True

    return str(
        valor
    ).strip() == ""


def formatear_numero(
    valor
):
    if valor_vacio(
        valor
    ):
        return ""

    try:
        numero = float(
            valor
        )

        if numero.is_integer():
            return str(
                int(
                    numero
                )
            )

        return str(
            round(
                numero,
                2
            )
        )

    except (
        TypeError,
        ValueError
    ):
        return str(
            valor
        )


def fuzz_ratio_seguro(
    texto_a,
    texto_b
):
    """
    Devuelve ambos puntajes usados para coincidencia de emergencia.
    """
    try:
        from rapidfuzz import fuzz

        return (
            fuzz.ratio(
                texto_a,
                texto_b
            ),
            fuzz.token_sort_ratio(
                texto_a,
                texto_b
            )
        )

    except Exception:
        return (
            0,
            0
        )


def normalizar(
    texto
):

    if texto is None:
        return ""

    texto = str(
        texto
    )

    texto = unicodedata.normalize(
        "NFD",
        texto
    )

    texto = "".join(
        caracter
        for caracter in texto
        if unicodedata.category(
            caracter
        ) != "Mn"
    )

    texto = texto.upper()

    texto = re.sub(
        r"[^A-Z0-9 ]",
        " ",
        texto
    )

    texto = re.sub(
        r"\s+",
        " ",
        texto
    )

    return texto.strip()


def limpiar_nombre_archivo(
    texto
):

    texto = normalizar(
        texto
    )

    texto = texto.replace(
        " ",
        "_"
    )

    if not texto:
        return "ALUMNO"

    return texto
